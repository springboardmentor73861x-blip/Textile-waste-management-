from collections import defaultdict
from datetime import datetime

from sqlalchemy.orm import Session

from app.models.inventory import Inventory
from app.models.upload import Upload
from app.models.report import Report


from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from openpyxl import Workbook



from app.services.environmental_impact_service import (
    get_environmental_impact,
)


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def _percentage(value):
    if value is None:
        return 0.0

    try:
        return float(
            str(value)
            .strip()
            .replace("%", "")
        )
    except (ValueError, TypeError):
        return 0.0


def _is_recyclable(value):
    if value is None:
        return False

    value = str(value).strip().lower()

    return (
        value in {
            "recyclable",
            "recycle",
            "recycled",
            "yes",
            "high",
            "excellent",
        }
        or "recyclable" in value
    )


def _month_label(value):
    return datetime.strptime(
        value,
        "%Y-%m"
    ).strftime("%b")


# =========================================================
# REPORTS & ANALYTICS DASHBOARD
# =========================================================

def get_report_analytics_dashboard(
    db: Session,
    company_code: str | None
):
    # -----------------------------------------------------
    # INVENTORY
    # -----------------------------------------------------

    inventory_items = (
    db.query(Inventory)
    .filter(
        Inventory.company_code == company_code
    )
    .order_by(
        Inventory.created_at.asc()
    )
    .all()
)

    # -----------------------------------------------------
    # TOTAL WASTE BATCHES
    # -----------------------------------------------------

    total_waste_batches = len(
        inventory_items
    )

    # -----------------------------------------------------
    # TOTAL QUANTITY / WEIGHT
    # -----------------------------------------------------

    total_quantity = sum(
        item.quantity or 0
        for item in inventory_items
    )

    # -----------------------------------------------------
    # RECYCLABLE QUANTITY
    # -----------------------------------------------------

    recyclable_quantity = sum(
        item.quantity or 0
        for item in inventory_items
        if _is_recyclable(
            item.recyclability
        )
    )

    if total_quantity > 0:
        recyclable_percentage = round(
            (
                recyclable_quantity
                / total_quantity
            ) * 100,
            2
        )
    else:
        recyclable_percentage = 0.0

    # -----------------------------------------------------
    # WASTE BY CATEGORY
    # -----------------------------------------------------

    waste_category_data = defaultdict(
        int
    )

    for item in inventory_items:

        category = (
            item.waste_category
            or "Unknown"
        )

        waste_category_data[
            category
        ] += item.quantity or 0

    waste_by_category = [
        {
            "category": category,
            "value": quantity
        }
        for category, quantity
        in waste_category_data.items()
    ]

    waste_by_category.sort(
        key=lambda item: item["value"],
        reverse=True
    )

    # -----------------------------------------------------
    # MATERIAL COMPOSITION
    # -----------------------------------------------------

    material_data = defaultdict(
        int
    )

    for item in inventory_items:

        material = (
            item.fabric_type
            or "Unknown"
        )

        material_data[
            material
        ] += item.quantity or 0

    material_composition = [
        {
            "material": material,
            "value": quantity
        }
        for material, quantity
        in material_data.items()
    ]

    material_composition.sort(
        key=lambda item: item["value"],
        reverse=True
    )

    # -----------------------------------------------------
    # RECYCLABILITY TREND
    # -----------------------------------------------------

    monthly_data = defaultdict(
        lambda: {
            "total": 0,
            "recyclable": 0
        }
    )

    for item in inventory_items:

        if not item.created_at:
            continue

        month = item.created_at.strftime(
            "%Y-%m"
        )

        quantity = item.quantity or 0

        monthly_data[
            month
        ]["total"] += quantity

        if _is_recyclable(
            item.recyclability
        ):
            monthly_data[
                month
            ]["recyclable"] += quantity

    months = sorted(
        monthly_data.keys()
    )[-6:]

    recyclability_trend = []

    for month in months:

        total = monthly_data[
            month
        ]["total"]

        recyclable = monthly_data[
            month
        ]["recyclable"]

        percentage = (
            (
                recyclable
                / total
            ) * 100
            if total > 0
            else 0
        )

        recyclability_trend.append(
            {
                "month": _month_label(
                    month
                ),
                "value": round(
                    percentage,
                    2
                )
            }
        )

    # -----------------------------------------------------
    # UPLOAD / ANALYSIS DATA
    # -----------------------------------------------------

    uploads = (
    db.query(Upload)
    .filter(
        Upload.company_code == company_code
    )
    .order_by(
        Upload.created_at.asc()
    )
    .all()
)

    # =====================================================
    # ENVIRONMENTAL IMPACT
    # =====================================================
    #
    # CO2, Water and Land are calculated from the
    # centralized environmental impact service.
    #
    # This means Reports & Analytics and Sustainability
    # Dashboard use the SAME calculation source.
    #
    # =====================================================

    environmental_impact = (
    get_environmental_impact(
        db,
        company_code
    )
)

    co2_saved = (
        environmental_impact[
            "co2_saved"
        ]
    )

    water_saved = (
        environmental_impact[
            "water_saved"
        ]
    )

    land_saved = (
        environmental_impact[
            "land_saved"
        ]
    )

    diverted_quantity = (
        environmental_impact[
            "diverted_quantity"
        ]
    )

    # -----------------------------------------------------
    # CO2 TREND
    # -----------------------------------------------------

    co2_trend = []

    upload_months = defaultdict(
        int
    )

    for upload in uploads:

        if not upload.created_at:
            continue

        month = upload.created_at.strftime(
            "%Y-%m"
        )

        upload_months[
            month
        ] += 1

    upload_month_list = sorted(
        upload_months.keys()
    )[-6:]

    # -----------------------------------------------------
    # IMPORTANT
    # -----------------------------------------------------
    # There is currently no historical environmental-impact
    # value stored per upload/month in the database.
    #
    # Therefore we do NOT create fake historical CO2 values.
    # The current total environmental impact comes from the
    # actual Inventory data.
    #
    # -----------------------------------------------------

    for month in upload_month_list:

        co2_trend.append(
            {
                "month": _month_label(
                    month
                ),
                "value": 0
            }
        )

    # -----------------------------------------------------
    # RECENT REPORTS
    # -----------------------------------------------------

    reports = (
    db.query(Report)
    .filter(
        Report.company_code == company_code
    )
    .order_by(
        Report.created_at.desc()
    )
    .limit(5)
    .all()
)

    recent_reports = []

    for report in reports:

        recent_reports.append(
            {
                "id": report.id,
                "name": report.report_name,
                "type": report.report_type,
                "status": report.status,
                "downloads": report.downloads or 0,
                "created_at": (
                    report.created_at.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                    if report.created_at
                    else "-"
                )
            }
        )

    # -----------------------------------------------------
    # SUMMARY
    # -----------------------------------------------------

    return {
        "stats": {

            # Existing data
            "total_waste_batches":
                total_waste_batches,

            "total_quantity":
                total_quantity,

            "recyclable_percentage":
                recyclable_percentage,

            # Environmental data
            "co2_saved":
                co2_saved,

            "water_saved":
                water_saved,

            "land_saved":
                land_saved,

            "diverted_quantity":
                diverted_quantity,

            "co2_unit":
                environmental_impact[
                    "co2_unit"
                ],

            "water_unit":
                environmental_impact[
                    "water_unit"
                ],

            "land_unit":
                environmental_impact[
                    "land_unit"
                ],

            "environmental_impact_estimated":
                environmental_impact[
                    "is_estimated"
                ],
        },

        "waste_by_category":
            waste_by_category,

        "material_composition":
            material_composition,

        "recyclability_trend":
            recyclability_trend,

        "co2_trend":
            co2_trend,

        "recent_reports":
            recent_reports
    }

def generate_analytics_report(
    db: Session,
    report_type: str,
    generated_by: str,
    company_name: str | None = None,
    company_code: str | None = None
):
    """
    Generate a real PDF report using the same PostgreSQL-backed
    dashboard data and existing environmental impact service.
    """

    # Get current dashboard data
    dashboard_data = get_report_analytics_dashboard(db, company_code)

    stats = dashboard_data["stats"]

    # Create Report database record
    report = Report(
    report_name=report_type,
    report_type=report_type,
    generated_by=generated_by,
    company_name=company_name,
    company_code=company_code,
    status="Completed",
    downloads=0
)

    db.add(report)
    db.commit()
    db.refresh(report)

    # Create PDF from dashboard data
    pdf_buffer = _create_report_pdf(
        dashboard_data=dashboard_data,
        report_type=report_type,
        generated_by=generated_by,
        report_id=report.id
    )

    return report, pdf_buffer

    # =========================================================
# CREATE PDF FROM DASHBOARD DATA
# =========================================================

def _create_report_pdf(
    dashboard_data,
    report_type,
    generated_by,
    report_id
):
    """
    Creates PDF from already calculated dashboard data.
    Does NOT create a database report record.
    """

    stats = dashboard_data["stats"]

    buffer = BytesIO()

    pdf = canvas.Canvas(
        buffer,
        pagesize=A4
    )

    width, height = A4

    y = height - 50

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    pdf.setFont(
        "Helvetica-Bold",
        18
    )

    pdf.drawString(
        50,
        y,
        report_type
    )

    y -= 30

    pdf.setFont(
        "Helvetica",
        10
    )

    pdf.drawString(
        50,
        y,
        f"Generated by: {generated_by}"
    )

    y -= 20

    pdf.drawString(
        50,
        y,
        f"Report ID: {report_id}"
    )

    y -= 35

    # -----------------------------------------------------
    # SUMMARY
    # -----------------------------------------------------

    pdf.setFont(
        "Helvetica-Bold",
        13
    )

    pdf.drawString(
        50,
        y,
        "Summary"
    )

    y -= 25

    pdf.setFont(
        "Helvetica",
        11
    )

    summary_lines = [
        f"Total Waste Batches: {stats['total_waste_batches']}",
        f"Total Quantity: {stats['total_quantity']}",
        f"Recyclable Percentage: {stats['recyclable_percentage']}%",
        f"CO2 Saved: {stats['co2_saved']} {stats['co2_unit']}",
        f"Water Saved: {stats['water_saved']} {stats['water_unit']}",
        f"Land Saved: {stats['land_saved']} {stats['land_unit']}",
        f"Diverted Quantity: {stats['diverted_quantity']}",
    ]

    for line in summary_lines:

        pdf.drawString(
            60,
            y,
            line
        )

        y -= 20

        if y < 60:

            pdf.showPage()

            y = height - 50

            pdf.setFont(
                "Helvetica",
                11
            )

    # -----------------------------------------------------
    # WASTE CLASSIFICATION
    # -----------------------------------------------------

    y -= 15

    pdf.setFont(
        "Helvetica-Bold",
        13
    )

    pdf.drawString(
        50,
        y,
        "Waste Classification"
    )

    y -= 25

    pdf.setFont(
        "Helvetica",
        10
    )

    for item in dashboard_data[
        "waste_by_category"
    ]:

        pdf.drawString(
            60,
            y,
            f"{item['category']}: {item['value']}"
        )

        y -= 18

        if y < 60:

            pdf.showPage()

            y = height - 50

            pdf.setFont(
                "Helvetica",
                10
            )

    # -----------------------------------------------------
    # MATERIAL COMPOSITION
    # -----------------------------------------------------

    y -= 15

    pdf.setFont(
        "Helvetica-Bold",
        13
    )

    pdf.drawString(
        50,
        y,
        "Material Composition"
    )

    y -= 25

    pdf.setFont(
        "Helvetica",
        10
    )

    for item in dashboard_data[
        "material_composition"
    ]:

        pdf.drawString(
            60,
            y,
            f"{item['material']}: {item['value']}"
        )

        y -= 18

        if y < 60:

            pdf.showPage()

            y = height - 50

            pdf.setFont(
                "Helvetica",
                10
            )

    # -----------------------------------------------------
    # RECYCLABILITY TREND
    # -----------------------------------------------------

    y -= 15

    pdf.setFont(
        "Helvetica-Bold",
        13
    )

    pdf.drawString(
        50,
        y,
        "Recyclability Trend"
    )

    y -= 25

    pdf.setFont(
        "Helvetica",
        10
    )

    for item in dashboard_data[
        "recyclability_trend"
    ]:

        pdf.drawString(
            60,
            y,
            f"{item['month']}: {item['value']}%"
        )

        y -= 18

        if y < 60:

            pdf.showPage()

            y = height - 50

            pdf.setFont(
                "Helvetica",
                10
            )

    # -----------------------------------------------------
    # ENVIRONMENTAL IMPACT
    # -----------------------------------------------------

    y -= 15

    if y < 80:

        pdf.showPage()

        y = height - 50

    pdf.setFont(
        "Helvetica-Bold",
        13
    )

    pdf.drawString(
        50,
        y,
        "Environmental Impact"
    )

    y -= 25

    pdf.setFont(
        "Helvetica",
        10
    )

    if stats[
        "environmental_impact_estimated"
    ]:

        pdf.drawString(
            60,
            y,
            "Environmental impact values are estimated."
        )

    else:

        pdf.drawString(
            60,
            y,
            "Environmental impact values are calculated from platform data."
        )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    pdf.setFont(
        "Helvetica",
        8
    )

    pdf.drawString(
        50,
        30,
        "Textile Waste Intelligence Platform"
    )

    pdf.save()

    buffer.seek(0)

    return buffer

# =========================================================
# EXPORT REPORTS & ANALYTICS TO EXCEL
# =========================================================

def export_analytics_excel(
    db: Session,
    company_code: str | None
):
    """
    Export current Reports & Analytics dashboard data
    into an Excel workbook using PostgreSQL-backed data.
    """

    dashboard_data = get_report_analytics_dashboard(
        db, company_code
    )

    stats = dashboard_data["stats"]

    workbook = Workbook()

    # =====================================================
    # SUMMARY SHEET
    # =====================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet.append([
        "Metric",
        "Value"
    ])

    summary_sheet.append([
        "Total Waste Batches",
        stats["total_waste_batches"]
    ])

    summary_sheet.append([
        "Total Quantity",
        stats["total_quantity"]
    ])

    summary_sheet.append([
        "Recyclable Percentage",
        stats["recyclable_percentage"]
    ])

    summary_sheet.append([
        "CO2 Saved",
        stats["co2_saved"]
    ])

    summary_sheet.append([
        "CO2 Unit",
        stats["co2_unit"]
    ])

    summary_sheet.append([
        "Water Saved",
        stats["water_saved"]
    ])

    summary_sheet.append([
        "Water Unit",
        stats["water_unit"]
    ])

    summary_sheet.append([
        "Land Saved",
        stats["land_saved"]
    ])

    summary_sheet.append([
        "Land Unit",
        stats["land_unit"]
    ])

    summary_sheet.append([
        "Diverted Quantity",
        stats["diverted_quantity"]
    ])

    summary_sheet.append([
        "Environmental Impact Estimated",
        stats["environmental_impact_estimated"]
    ])

    # =====================================================
    # WASTE CLASSIFICATION SHEET
    # =====================================================

    waste_sheet = workbook.create_sheet(
        "Waste Classification"
    )

    waste_sheet.append([
        "Waste Category",
        "Quantity"
    ])

    for item in dashboard_data[
        "waste_by_category"
    ]:

        waste_sheet.append([
            item["category"],
            item["value"]
        ])

    # =====================================================
    # MATERIAL COMPOSITION SHEET
    # =====================================================

    material_sheet = workbook.create_sheet(
        "Material Composition"
    )

    material_sheet.append([
        "Material",
        "Quantity"
    ])

    for item in dashboard_data[
        "material_composition"
    ]:

        material_sheet.append([
            item["material"],
            item["value"]
        ])

    # =====================================================
    # RECYCLABILITY TREND SHEET
    # =====================================================

    recyclability_sheet = workbook.create_sheet(
        "Recyclability Trend"
    )

    recyclability_sheet.append([
        "Month",
        "Recyclability Percentage"
    ])

    for item in dashboard_data[
        "recyclability_trend"
    ]:

        recyclability_sheet.append([
            item["month"],
            item["value"]
        ])

    # =====================================================
    # CO2 TREND SHEET
    # =====================================================

    co2_sheet = workbook.create_sheet(
        "CO2 Trend"
    )

    co2_sheet.append([
        "Month",
        "CO2 Trend Value"
    ])

    for item in dashboard_data[
        "co2_trend"
    ]:

        co2_sheet.append([
            item["month"],
            item["value"]
        ])

    # =====================================================
    # RECENT REPORTS SHEET
    # =====================================================

    reports_sheet = workbook.create_sheet(
        "Recent Reports"
    )

    reports_sheet.append([
        "Report ID",
        "Report Name",
        "Report Type",
        "Status",
        "Downloads",
        "Created At"
    ])

    for report in dashboard_data[
        "recent_reports"
    ]:

        reports_sheet.append([
            report["id"],
            report["name"],
            report["type"],
            report["status"],
            report["downloads"],
            report["created_at"]
        ])

    # =====================================================
    # AUTO COLUMN WIDTH
    # =====================================================

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            max_length = 0

            column_letter = column[0].column_letter

            for cell in column:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                40
            )

    # =====================================================
    # SAVE WORKBOOK TO MEMORY
    # =====================================================

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer    